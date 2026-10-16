# processing/docx_reader.py
"""Word (.docx) 結構化讀取：按文件順序輸出段落與表格，還原標題 / 自動編號 / 文字方塊 / 上下標，並分頁。

移植並精簡自 word.py 的 _parse_docx / _docx_num_prefix / _iter_docx_block_items。
僅在「沒有 LibreOffice」時使用（有 LibreOffice 時 docx 會轉 PDF 走版面還原，得到真實分頁與圖片）。
"""
import re

from processing.textutil import clean_inline, cn_num, restore_marks, rows_to_markdown_table, SUP_O, SUP_C, SUB_O, SUB_C

_LEAD_NUM = re.compile(r"^\s*[（(]?\s*[\dA-Za-z一二三四五六七八九十①-⑩]+\s*[)）.、．]")


def _iter_blocks(doc):
    """依文件順序產生 Paragraph / Table（python-docx 預設把兩者分開，順序會丟失）。"""
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


def _num_prefix(doc, para, counters):
    """還原 Word 自動編號（題號 / 選項），失敗回傳空字串。"""
    try:
        num_pr = para._element.xpath("./w:pPr/w:numPr")
        if not num_pr:
            return ""
        num_id = num_pr[0].xpath("./w:numId/@w:val")
        ilvl = (num_pr[0].xpath("./w:ilvl/@w:val") or ["0"])[0]
        if not num_id or num_id[0] == "0":
            return ""
        num_id = num_id[0]
        numbering = doc.part.numbering_part.element
        aid = numbering.xpath(f'./w:num[@w:numId="{num_id}"]/w:abstractNumId/@w:val')
        fmt = numbering.xpath(f'./w:abstractNum[@w:abstractNumId="{aid[0]}"]/w:lvl[@w:ilvl="{ilvl}"]/w:numFmt/@w:val')[0]
        key = (num_id, int(ilvl))
        counters[key] = counters.get(key, 0) + 1
        for k in list(counters):                       # 子層級在上層遞增時歸零
            if k[0] == num_id and k[1] > int(ilvl):
                counters[k] = 0
        n = counters[key]
        if fmt == "decimal":
            return f"{n}. "
        if fmt == "decimalEnclosedParen":
            return f"({n}) "
        if fmt in ("upperLetter", "lowerLetter"):
            return f"({chr(64 + n)}) " if n <= 26 else ""
        if fmt in ("chineseCounting", "chineseCountingThousand", "ideographTraditional", "japaneseCounting"):
            return f"{cn_num(n)}、"
        if fmt == "decimalEnclosedCircle" and n <= 10:
            return "①②③④⑤⑥⑦⑧⑨⑩"[n - 1] + " "
    except Exception:
        pass
    return ""


def _paragraph_text(p):
    """段落文字；若含上 / 下標 run，改由 run 組裝以保留 <sup>/<sub>。"""
    try:
        runs = list(p.runs)
        if any(r.font.superscript or r.font.subscript for r in runs):
            out = []
            for r in runs:
                t = r.text
                if t and r.font.superscript:
                    t = SUP_O + t + SUP_C
                elif t and r.font.subscript:
                    t = SUB_O + t + SUB_C
                out.append(t)
            return "".join(out)
    except Exception:
        pass
    return p.text


def _table_rows(tbl):
    rows = []
    for r in tbl.rows:
        cells = []
        for c in r.cells:
            val = c.text.strip()
            if not cells or cells[-1] != val:          # 合併儲存格會重複出現，去掉相鄰重複
                cells.append(val)
        rows.append(cells)
    return rows


def read_docx_blocks(path, normalize=True):
    """回傳 [(kind, text)]，kind: 'h' 標題 / 'p' 段落 / 'q' 文字方塊 / 't' 表格(Markdown) / 'pb' 明確分頁。"""
    import docx
    from docx.table import Table
    doc = docx.Document(path)
    norm = clean_inline if normalize else (lambda s: s.strip())
    blocks, counters, seen_hdr = [], {}, set()

    for section in doc.sections:                       # 頁眉（去重）
        try:
            for p in section.header.paragraphs:
                t = p.text.strip()
                if t and t not in seen_hdr:
                    seen_hdr.add(t)
                    blocks.append(("h1", norm(t)))
            for tbl in section.header.tables:
                md = rows_to_markdown_table(_table_rows(tbl), norm)
                if md:
                    blocks.append(("t", md))
        except Exception:
            pass

    for block in _iter_blocks(doc):
        if isinstance(block, Table):
            md = rows_to_markdown_table(_table_rows(block), norm)
            if md:
                blocks.append(("t", md))
            continue
        txbx = [x.strip() for x in block._element.xpath(".//w:txbxContent//w:t/text()") if x.strip()]
        if txbx:
            blocks.append(("q", norm(" ".join(txbx))))
        level = 0
        try:
            sname = block.style.name or ""
            m = re.search(r"(?:Heading|標題)\s*(\d)", sname)
            if m:
                level = int(m.group(1))
            elif sname.lower() == "title":
                level = 1
        except Exception:
            pass
        text = _paragraph_text(block)
        if text.strip():
            lines = [s for s in text.split("\n") if s.strip()]
            prefix = _num_prefix(doc, block, counters)
            if prefix and lines and not _LEAD_NUM.match(lines[0]):
                lines[0] = prefix + lines[0].lstrip()
            for i, s in enumerate(lines):
                blocks.append((f"h{level}" if (level and i == 0) else "p", norm(s)))
        if block._element.xpath('.//w:br[@w:type="page"]'):
            blocks.append(("pb", ""))
    return blocks


def read_docx_pages(path, page_chars=1200, normalize=True):
    """依明確分頁符與字數把內容切成「頁」，回傳每頁文字（至少一頁）。表格不會被切開。"""
    pages, cur, cur_len = [], [], 0

    def render(items):
        out = []
        for kind, text in items:
            if kind == "t":
                if out:
                    out.append("")
                out.append(restore_marks(text))
                out.append("")
            elif kind == "q":
                out.append("> " + restore_marks(text))
            elif kind.startswith("h") and kind != "h0":
                out.append("#" * min(int(kind[1:]), 6) + " " + restore_marks(text))
            else:
                out.append(restore_marks(text))
        return "\n".join(out).strip()

    def close():
        nonlocal cur, cur_len
        if cur:
            pages.append(render(cur))
        cur, cur_len = [], 0

    for kind, text in read_docx_blocks(path, normalize):
        if kind == "pb":
            close()
            continue
        if cur_len >= page_chars and kind != "t" or (kind == "t" and cur_len + len(text) > page_chars * 1.5 and cur):
            close()
        cur.append((kind, text))
        cur_len += len(text)
    close()
    return pages or [""]
